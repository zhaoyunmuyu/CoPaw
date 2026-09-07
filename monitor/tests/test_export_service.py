# -*- coding: utf-8 -*-
"""Tests for Monitor cron export service."""

import pytest
from datetime import datetime, timezone
from io import BytesIO

pytest.importorskip("openpyxl")

from monitor.app.models.cron import CronJobModel, ExecutionModel
from monitor.app.services.cron.export_service import ExportService


class TestExportService:
    """Tests for ExportService."""

    @pytest.fixture
    def export_service(self):
        """Create an ExportService instance."""
        return ExportService()

    @pytest.fixture
    def sample_jobs(self):
        """Create sample CronJobModel instances."""
        return [
            CronJobModel(
                id="job-001",
                name="Daily Report",
                tenant_id="tenant-001",
                enabled=True,
                task_type="agent",
                cron_expr="0 9 * * *",
                timezone="Asia/Shanghai",
                channel="console",
                creator_user_id="user-001",
                created_at=datetime.now(timezone.utc),
                updated_at=datetime.now(timezone.utc),
            ),
            CronJobModel(
                id="job-002",
                name="Weekly Summary",
                tenant_id="tenant-002",
                enabled=False,
                task_type="text",
                cron_expr="0 10 * * mon",
                timezone="UTC",
                channel="zhaohu",
                status="paused",
                creator_user_id="user-002",
                created_at=datetime.now(timezone.utc),
                updated_at=datetime.now(timezone.utc),
            ),
        ]

    @pytest.fixture
    def sample_executions(self):
        """Create sample ExecutionModel instances."""
        now = datetime.now(timezone.utc)
        return [
            ExecutionModel(
                id=1,
                job_id="job-001",
                job_name="Daily Report",
                tenant_id="tenant-001",
                actual_time=now,
                end_time=now,
                status="success",
                duration_ms=1500,
                is_manual=False,
                trace_id="trace-001",
                session_id="session-001",
                output_preview="Task completed successfully",
            ),
            ExecutionModel(
                id=2,
                job_id="job-001",
                job_name="Daily Report",
                tenant_id="tenant-001",
                actual_time=now,
                end_time=now,
                status="error",
                duration_ms=500,
                error_message="Timeout exceeded",
                is_manual=True,
            ),
        ]

    def test_export_jobs_returns_bytes(self, export_service, sample_jobs):
        """Test that export_jobs returns bytes."""
        result = export_service.export_jobs(sample_jobs)
        assert isinstance(result, bytes)
        assert len(result) > 0

    def test_export_executions_returns_bytes(
        self,
        export_service,
        sample_executions,
    ):
        """Test that export_executions returns bytes."""
        result = export_service.export_executions(sample_executions)
        assert isinstance(result, bytes)
        assert len(result) > 0

    def test_export_jobs_empty_list(self, export_service):
        """Test export_jobs with empty list."""
        result = export_service.export_jobs([])
        assert isinstance(result, bytes)
        assert len(result) > 0  # Excel file has headers even with no data

    def test_export_executions_empty_list(self, export_service):
        """Test export_executions with empty list."""
        result = export_service.export_executions([])
        assert isinstance(result, bytes)
        assert len(result) > 0

    def test_export_jobs_has_correct_columns(
        self,
        export_service,
        sample_jobs,
    ):
        """Test that export_jobs Excel has correct structure."""
        result = export_service.export_jobs(sample_jobs)
        # Verify it's a valid Excel file by checking the bytes
        # Excel files start with specific bytes (PK for ZIP-based format)
        assert result[:2] == b"PK"  # Excel uses ZIP container

    def test_export_skill_usage_details_masks_names_and_status(
        self,
        export_service,
    ):
        result = export_service.export_skill_usage_details(
            [
                {
                    "task_name": "任务",
                    "tenant_id": "u1",
                    "tenant_name": "经理",
                    "bbk_id": "001",
                    "task_status": "active",
                    "execution_id": 7,
                    "actual_time": datetime(2026, 9, 3, 10, 0, 0),
                    "execution_status": "success",
                    "async_status": "error",
                    "is_read": 1,
                    "custuid": "c1",
                    "cust_nm": "张三丰",
                    "clicked_plan": 1,
                    "clicked_phone": 0,
                    "clicked_insight": 1,
                },
            ],
        )
        from openpyxl import load_workbook

        sheet = load_workbook(BytesIO(result), read_only=True).active
        values = list(sheet.iter_rows(min_row=2, values_only=True))[0]
        assert values[5] == 7
        assert values[7] == "子任务执行失败"
        assert values[10] == "张*丰"
        assert values[11:14] == ("是", "否", "是")
